from bs4 import BeautifulSoup
from pptx import Presentation

# Carrega o arquivo HTML
with open("devsecops_para_leigos v3.html", "r", encoding="utf-8") as f:
    soup = BeautifulSoup(f, "html.parser")

# Inicia a apresentação do PowerPoint
prs = Presentation()

# Encontra todos os slides do HTML
slides_html = soup.find_all("section", class_="slide-container")

for slide_html in slides_html:
    # Usa o layout padrão de "Título e Conteúdo"
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    # Extrai o título (procura por h1 ou h2)
    title_element = slide_html.find(["h1", "h2"])
    if title_element:
        slide.shapes.title.text = title_element.get_text(strip=True)
        
    # Seleciona a caixa de texto principal do slide
    content_element = slide.placeholders[1]
    text_frame = content_element.text_frame
    text_frame.text = ""
    
    # Extrai o conteúdo (parágrafos, listas, subtítulos e citações)
    for element in slide_html.find_all(["p", "ul", "h3", "h4", "blockquote"]):
        # Se for uma lista, adiciona os marcadores adequadamente
        if element.name == "ul":
            for li in element.find_all("li"):
                p = text_frame.add_paragraph()
                p.text = "• " + li.get_text(strip=True)
                p.level = 1
        else:
            p = text_frame.add_paragraph()
            p.text = element.get_text(strip=True)
            p.level = 0
            
# Salva o arquivo final
prs.save("DevSecOps_Apresentacao.pptx")
print("Arquivo DevSecOps_Apresentacao.pptx gerado com sucesso!")